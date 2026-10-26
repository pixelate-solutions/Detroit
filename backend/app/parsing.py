from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

import fitz
import pytesseract
from docx import Document
from pdf2image import convert_from_path


@dataclass
class ParsedDocument:
    full_text: str
    pages: list[str]
    ocr_used: bool


@dataclass
class TextChunk:
    text: str
    page_number: int
    char_start: int
    char_end: int


def extract_document(path: Path) -> ParsedDocument:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".docx":
        doc = Document(path)
        text = "\n".join(paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip())
        return ParsedDocument(full_text=text, pages=[text], ocr_used=False)
    if suffix in {".txt", ".md"}:
        text = path.read_text(encoding="utf-8", errors="replace")
        return ParsedDocument(full_text=text, pages=[text], ocr_used=False)
    raise ValueError("Supported file types are PDF, DOCX, TXT, and MD.")


def _extract_pdf(path: Path) -> ParsedDocument:
    pdf = fitz.open(path)
    pages: list[str] = []
    ocr_used = False
    for page_number, page in enumerate(pdf, start=1):
        text = page.get_text("text").strip()
        if len(re.sub(r"\s+", "", text)) < 40:
            try:
                image = convert_from_path(
                    str(path), dpi=240, first_page=page_number, last_page=page_number, fmt="png"
                )[0]
                text = pytesseract.image_to_string(image).strip()
                ocr_used = True
            except Exception as exc:
                raise RuntimeError(
                    f"Page {page_number} needs OCR, but local Tesseract/Poppler OCR failed: {exc}"
                ) from exc
        pages.append(text)
    pdf.close()
    return ParsedDocument(full_text="\n\n".join(pages), pages=pages, ocr_used=ocr_used)


def chunk_document(parsed: ParsedDocument, target_chars: int = 2600, overlap_chars: int = 360) -> list[TextChunk]:
    chunks: list[TextChunk] = []
    global_offset = 0
    for page_number, page_text in enumerate(parsed.pages, start=1):
        clean = re.sub(r"[ \t]+", " ", page_text).strip()
        if not clean:
            global_offset += 2
            continue
        start = 0
        while start < len(clean):
            end = min(len(clean), start + target_chars)
            if end < len(clean):
                boundary = max(clean.rfind("\n", start, end), clean.rfind(". ", start, end))
                if boundary > start + target_chars // 2:
                    end = boundary + 1
            text = clean[start:end].strip()
            if text:
                chunks.append(
                    TextChunk(
                        text=text,
                        page_number=page_number,
                        char_start=global_offset + start,
                        char_end=global_offset + end,
                    )
                )
            if end >= len(clean):
                break
            start = max(start + 1, end - overlap_chars)
        global_offset += len(clean) + 2
    return chunks


def infer_title_and_type(filename: str, text: str) -> tuple[str, str]:
    leading_lines = [line.strip() for line in text.splitlines() if line.strip()]
    first_line = leading_lines[0] if leading_lines else ""
    title_lines = [first_line] if first_line else []
    if first_line.isupper():
        for line in leading_lines[1:3]:
            if line.isupper() and len(" ".join(title_lines + [line])) <= 140:
                title_lines.append(line)
            else:
                break
    candidate_title = " ".join(title_lines)
    title = candidate_title if 4 <= len(candidate_title) <= 140 else Path(filename).stem.replace("_", " ").title()
    types = [
        ("meeting_minutes", ("minutes", "city council meeting")),
        ("correspondence", ("memorandum", "email", "correspondence", "complaint")),
        ("inspection_report", ("inspection", "code enforcement", "followup_ce")),
        ("permit", ("permit", "application")),
        ("contract", ("contract", "agreement")),
        ("ordinance", ("ordinance", "zoning amendment")),
    ]
    primary = f"{filename} {title}".lower()
    doc_type = next((kind for kind, markers in types if any(marker in primary for marker in markers)), None)
    if doc_type is None:
        sample = text[:700].lower()
        doc_type = next((kind for kind, markers in types if any(marker in sample for marker in markers)), "other")
    return title, doc_type
